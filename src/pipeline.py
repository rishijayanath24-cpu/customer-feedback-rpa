"""End-to-end RPA pipeline: read -> validate -> classify -> route -> report."""
from __future__ import annotations

import csv
import json
import logging
import re
import shutil
from collections import Counter
from datetime import datetime, timedelta
from pathlib import Path

from classifier import FeedbackClassifier

log = logging.getLogger("feedback_rpa")

REQUIRED_COLUMNS = ["feedback_id", "date", "customer_name", "email", "channel", "rating", "feedback_text"]
OUT_COLUMNS = REQUIRED_COLUMNS + [
    "category", "department", "sentiment", "priority", "sla_due",
    "matched_keywords", "status", "notes",
]
EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")
PRIORITY_ORDER = {"Urgent": 0, "High": 1, "Normal": 2, "Low": 3}


def read_feedback(paths: Path | list[Path]) -> list[dict]:
    """Read one or more feedback CSVs. Several files let the web app's live submissions
    (data/feedback_live.csv) be processed alongside the batch export."""
    if isinstance(paths, (str, Path)):
        paths = [Path(paths)]
    rows = []
    for path in paths:
        with open(path, newline="", encoding="utf-8-sig") as fh:
            reader = csv.DictReader(fh)
            missing = [c for c in REQUIRED_COLUMNS if c not in (reader.fieldnames or [])]
            if missing:
                raise ValueError(f"{path.name} is missing required columns: {missing}")
            rows += [{k: (v or "").strip() for k, v in row.items()} for row in reader]
    return rows


def parse_rating(raw: str) -> int | None:
    try:
        value = int(float(raw))
    except ValueError:
        return None
    return value if 1 <= value <= 5 else None


def validate(rows: list[dict]) -> tuple[list[dict], list[dict]]:
    """Split into (valid, exceptions). Exceptions get a 'notes' explaining why."""
    valid, exceptions, seen = [], [], set()
    for row in rows:
        problems = []
        if not row["feedback_id"]:
            problems.append("missing feedback_id")
        elif row["feedback_id"] in seen:
            problems.append("duplicate feedback_id")
        if not row["feedback_text"]:
            problems.append("empty feedback text")
        if row["email"] and not EMAIL_RE.match(row["email"]):
            problems.append("invalid email")
        if row["feedback_id"]:
            seen.add(row["feedback_id"])
        if problems:
            exceptions.append({**row, "status": "Exception", "notes": "; ".join(problems)})
        else:
            valid.append(row)
    return valid, exceptions


def process(rows: list[dict], clf: FeedbackClassifier, sla_hours: dict, now: datetime) -> list[dict]:
    out = []
    for row in rows:
        try:
            result = clf.classify(row["feedback_text"], parse_rating(row["rating"]))
        except Exception as exc:  # one bad record must never stop the whole batch
            log.exception("Classification failed for %s", row["feedback_id"])
            out.append({**row, "status": "Exception", "notes": f"classification error: {exc}"})
            continue
        due = now + timedelta(hours=sla_hours[result.priority])
        out.append({
            **row,
            "category": result.category,
            "department": result.department,
            "sentiment": result.sentiment,
            "priority": result.priority,
            "sla_due": due.strftime("%Y-%m-%d %H:%M"),
            "matched_keywords": ", ".join(result.matched_keywords),
            "status": "Needs Review" if result.needs_review else "Routed",
            "notes": result.reason,
        })
    return out


def _safe_name(name: str) -> str:
    return re.sub(r"[^A-Za-z0-9]+", "_", name).strip("_")


def write_csv(path: Path, rows: list[dict]) -> None:
    with open(path, "w", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(fh, fieldnames=OUT_COLUMNS, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def route(rows: list[dict], out_dir: Path) -> dict[str, list[dict]]:
    """Write one queue file per department (most urgent first) plus a notification."""
    by_dept: dict[str, list[dict]] = {}
    for row in rows:
        by_dept.setdefault(row["department"], []).append(row)

    dept_dir, note_dir = out_dir / "by_department", out_dir / "notifications"
    for d in (dept_dir, note_dir):
        shutil.rmtree(d, ignore_errors=True)
        d.mkdir(parents=True)

    for dept, items in by_dept.items():
        items.sort(key=lambda r: (PRIORITY_ORDER[r["priority"]], r["date"]))
        write_csv(dept_dir / f"{_safe_name(dept)}.csv", items)
        counts = Counter(i["priority"] for i in items)
        lines = [
            f"To: {_safe_name(dept).lower()}@company.example",
            f"Subject: [Feedback Routing] {len(items)} new item(s) for {dept}"
            + (f" - {counts['Urgent']} URGENT" if counts["Urgent"] else ""),
            "",
            f"Hello {dept} team,",
            "",
            f"{len(items)} customer feedback record(s) were routed to you. Full details: "
            f"by_department/{_safe_name(dept)}.csv",
            "",
        ]
        for i in items:
            lines.append(
                f"  [{i['priority']:<6}] {i['feedback_id']} | {i['category']} | "
                f"due {i['sla_due']} | {i['customer_name']}"
            )
        (note_dir / f"{_safe_name(dept)}.txt").write_text("\n".join(lines) + "\n", encoding="utf-8")
    return by_dept


def write_excel(path: Path, all_rows, exceptions, by_dept, summary) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    fills = {"Urgent": "F8B4B4", "High": "FBD9A6", "Normal": "FFFFFF", "Low": "D5EBD0"}
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append(["Customer Feedback Classification & Routing - Run Summary"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([])
    for label, value in [
        ("Run time", summary["run_time"]),
        ("Records read", summary["records_read"]),
        ("Routed automatically", summary["routed"]),
        ("Needs manual review", summary["needs_review"]),
        ("Exceptions (invalid)", summary["exceptions"]),
    ]:
        ws.append([label, value])
    for title, key in [("By department", "by_department"), ("By category", "by_category"),
                       ("By priority", "by_priority"), ("By sentiment", "by_sentiment")]:
        ws.append([])
        ws.append([title, "Count"])
        ws.cell(ws.max_row, 1).font = ws.cell(ws.max_row, 2).font = Font(bold=True)
        for name, n in sorted(summary[key].items(), key=lambda kv: -kv[1]):
            ws.append([name, n])
    ws.column_dimensions["A"].width = 32

    def sheet(title: str, rows: list[dict]) -> None:
        s = wb.create_sheet(title[:31])
        s.append(OUT_COLUMNS)
        for c in s[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="305496")
        for r in rows:
            s.append([r.get(c, "") for c in OUT_COLUMNS])
            fill = fills.get(r.get("priority", ""))
            if fill:
                s.cell(s.max_row, OUT_COLUMNS.index("priority") + 1).fill = PatternFill("solid", fgColor=fill)
        for idx, col in enumerate(OUT_COLUMNS, 1):
            width = 60 if col in ("feedback_text", "notes") else max(12, len(col) + 2)
            s.column_dimensions[get_column_letter(idx)].width = width
        s.freeze_panes = "A2"
        s.auto_filter.ref = s.dimensions

    sheet("All Feedback", all_rows)
    for dept, items in sorted(by_dept.items()):
        sheet(dept.replace("&", "and"), items)
    if exceptions:
        sheet("Exceptions", exceptions)
    wb.save(path)


def run(input_path: Path | list[Path], rules_path: Path, out_dir: Path,
        now: datetime | None = None) -> dict:
    now = now or datetime.now()
    out_dir.mkdir(parents=True, exist_ok=True)
    clf = FeedbackClassifier.from_file(rules_path)
    sla = clf.rules["sla_hours"]

    log.info("Reading %s", input_path)
    rows = read_feedback(input_path)
    log.info("Read %d records", len(rows))

    valid, exceptions = validate(rows)
    log.info("Validation: %d valid, %d exceptions", len(valid), len(exceptions))
    for e in exceptions:
        log.warning("Exception %s: %s", e["feedback_id"] or "<no id>", e["notes"])

    processed = process(valid, clf, sla, now)
    # classification failures from process() join the exception queue
    exceptions += [r for r in processed if r["status"] == "Exception"]
    processed = [r for r in processed if r["status"] != "Exception"]

    by_dept = route(processed, out_dir)
    for dept, items in sorted(by_dept.items()):
        log.info("Routed %2d record(s) -> %s", len(items), dept)

    write_csv(out_dir / "classified_feedback.csv", processed)
    if exceptions:
        write_csv(out_dir / "exceptions.csv", exceptions)
    elif (out_dir / "exceptions.csv").exists():
        (out_dir / "exceptions.csv").unlink()

    summary = {
        "run_time": now.strftime("%Y-%m-%d %H:%M:%S"),
        "records_read": len(rows),
        "routed": sum(r["status"] == "Routed" for r in processed),
        "needs_review": sum(r["status"] == "Needs Review" for r in processed),
        "exceptions": len(exceptions),
        "by_department": dict(Counter(r["department"] for r in processed)),
        "by_category": dict(Counter(r["category"] for r in processed)),
        "by_priority": dict(Counter(r["priority"] for r in processed)),
        "by_sentiment": dict(Counter(r["sentiment"] for r in processed)),
    }
    (out_dir / "summary.json").write_text(json.dumps(summary, indent=2), encoding="utf-8")
    write_excel(out_dir / "feedback_report.xlsx", processed, exceptions, by_dept, summary)
    log.info("Done: %s routed, %s need review, %s exceptions",
             summary["routed"], summary["needs_review"], summary["exceptions"])
    return summary
